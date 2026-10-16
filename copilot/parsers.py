"""Document parsers: turn PDF / DOCX / TXT RFPs into clean plain text."""
import os
import re
import subprocess


def _clean(text):
    text = text.replace("\r\n", "\n").replace("\x0c", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)          # collapse blank runs
    lines = [ln.rstrip() for ln in text.split("\n")]
    # drop lines that are just page numbers
    lines = [ln for ln in lines if not re.fullmatch(r"\s*(page\s+)?\d+(\s+of\s+\d+)?\s*", ln, re.I)]
    return "\n".join(lines).strip()


def _pdf_text(path):
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except ImportError:
        pass
    try:
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(path).pages)
    except ImportError:
        pass
    # last resort: poppler's pdftotext
    r = subprocess.run(["pdftotext", "-layout", path, "-"],
                       capture_output=True, text=True)
    if r.returncode == 0:
        return r.stdout
    raise RuntimeError("No PDF parser available (install pdfplumber or pypdf, "
                       "or ensure pdftotext is on PATH).")


def _docx_text(path):
    from docx import Document  # python-docx
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_text(path):
    """Return cleaned plain text for a .pdf, .docx, or .txt file."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        raw = _pdf_text(path)
    elif ext in (".docx", ".doc"):
        if ext == ".doc":
            raise RuntimeError("Legacy .doc not supported; save as .docx first.")
        raw = _docx_text(path)
    elif ext in (".txt", ".md", ""):
        with open(path, encoding="utf-8", errors="replace") as f:
            raw = f.read()
    else:
        raise RuntimeError(f"Unsupported file type: {ext}")
    text = _clean(raw)
    if len(text) < 100:
        raise RuntimeError(f"Extracted only {len(text)} characters from {path} - "
                           "the file may be scanned images (needs OCR) or empty.")
    return text
